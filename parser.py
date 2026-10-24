# -*- coding: utf-8 -*-
"""
TRỘN ĐỀ PRO (SMART EXAM SHUFFLER & MULTI-VERSION GENERATOR)
Module: parser.py - Trích xuất câu hỏi từ file Word DOCX với bảo toàn XML MathType/OMML/Hình ảnh
"""

import re
import copy
from typing import List, Tuple, Optional
import docx
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn

from schema import Question, Option, PartType, CognitiveLevel


DEPENDENT_PHRASES = [
    r"cả\s+[a-d]\s+(và|với)\s+[a-d]",
    r"cả\s+hai\s+phương\s+án",
    r"tất\s+cả\s+(các\s+)?đáp\s+án\s+trên",
    r"không\s+có\s+(đáp\s+án|phương\s+án)\s+nào",
    r"các\s+phương\s+án\s+trên\s+đều\s+đúng",
    r"cả\s+3\s+đáp\s+án\s+trên",
]


class ExamDocxParser:
    def __init__(self, docx_path: str):
        self.docx_path = docx_path
        self.doc = docx.Document(docx_path)
        self.questions: List[Question] = []
        self.warnings: List[str] = []

    def parse(self) -> List[Question]:
        current_part = PartType.PART_1_CHOICE
        current_group_id = None
        raw_paragraphs = self.doc.paragraphs

        # Danh sách gom nhóm các paragraph cho từng câu
        question_blocks = []
        current_block = []

        part_pattern = re.compile(r"^\s*PHẦN\s+([I|V|X|1-4]+)", re.IGNORECASE)
        q_start_pattern = re.compile(r"^\s*Câu\s+(\d+)(?:\s*\([^\)]+\))?[\.:\s]", re.IGNORECASE)
        group_start_pattern = re.compile(r"\[NHÓM\]", re.IGNORECASE)
        group_end_pattern = re.compile(r"\[HẾT NHÓM\]", re.IGNORECASE)

        def _is_end_of_exam(text: str) -> bool:
            s = text.strip()
            if not s:
                return False
            if re.search(r"HẾT\s+PHẦN", s, re.IGNORECASE):
                return False
            if re.search(r"[-=_*~]{2,}\s*HẾT\b", s, re.IGNORECASE) or re.search(r"\bHẾT\s*[-=_*~]{2,}", s, re.IGNORECASE):
                return True
            if re.match(r"^[-=_*~]*\s*HẾT\s*[-=_*~]*$", s, re.IGNORECASE):
                return True
            if re.match(r"^(PHẦN\s+)?(ĐÁP ÁN|HƯỚNG DẪN CHẤM|HƯỚNG DẪN GIẢI|BẢNG ĐÁP ÁN)(\s*VÀ\s*THANG\s*ĐIỂM)?\s*$", s, re.IGNORECASE):
                return True
            return False

        answer_paragraphs = []
        for p_idx, p in enumerate(raw_paragraphs):
            text = p.text.strip()
            if not text:
                continue

            # Kiểm tra kết thúc đề thi / bắt đầu phần đáp án
            if (len(question_blocks) > 0 or len(current_block) > 0) and _is_end_of_exam(text):
                if current_block:
                    question_blocks.append(current_block)
                    current_block = []
                answer_paragraphs = raw_paragraphs[p_idx:]
                break

            # Kiểm tra chuyển phần
            part_match = part_pattern.match(text)
            if part_match:
                p_sym = part_match.group(1).upper()
                if p_sym in ["I", "1"]:
                    current_part = PartType.PART_1_CHOICE
                elif p_sym in ["II", "2"]:
                    current_part = PartType.PART_2_TRUE_FALSE
                elif p_sym in ["III", "3"]:
                    current_part = PartType.PART_3_SHORT_ANSWER
                elif p_sym in ["IV", "4"]:
                    current_part = PartType.PART_4_ESSAY
                continue

            # Kiểm tra nhóm câu chùm
            if group_start_pattern.search(text):
                current_group_id = f"GROUP_{len(question_blocks) + 1}"
                continue
            if group_end_pattern.search(text):
                current_group_id = None
                continue

            # Kiểm tra bắt đầu câu mới
            q_match = q_start_pattern.match(text)
            if q_match:
                if current_block:
                    question_blocks.append(current_block)
                current_block = [(p, current_part, current_group_id)]
            else:
                if current_block:
                    current_block.append((p, current_part, current_group_id))

        if current_block:
            question_blocks.append(current_block)

        # Xử lý từng khối câu hỏi
        for idx, block in enumerate(question_blocks, 1):
            q = self._process_question_block(block, idx)
            if q:
                self.questions.append(q)

        # Trích xuất lời giải tự luận từ phần đáp án sau đề thi (nếu có)
        if answer_paragraphs:
            current_ans_q_num = None
            ans_paragraphs_by_q = {}
            ans_q_pattern = re.compile(r"^\s*(?:Câu|Bài)\s+(\d+)[\.:\s]*", re.IGNORECASE)

            for p in answer_paragraphs:
                txt = p.text.strip()
                if not txt:
                    continue
                m = ans_q_pattern.match(txt)
                if m:
                    current_ans_q_num = int(m.group(1))
                    rest = txt[m.end():].strip()
                    if current_ans_q_num not in ans_paragraphs_by_q:
                        ans_paragraphs_by_q[current_ans_q_num] = []
                    if rest:
                        ans_paragraphs_by_q[current_ans_q_num].append(p)
                else:
                    if current_ans_q_num and not re.match(r"^(Phần\s+\d+|ĐÁP ÁN|HƯỚNG DẪN|BẢNG ĐÁP ÁN)", txt, re.IGNORECASE):
                        ans_paragraphs_by_q[current_ans_q_num].append(p)

            # Gán lời giải cho các câu hỏi tương ứng
            for q in self.questions:
                q_num = q.original_index
                if q_num in ans_paragraphs_by_q and not q.solution_xml_elements:
                    q.solution_xml_elements = [copy.deepcopy(p._p) for p in ans_paragraphs_by_q[q_num]]

        # Đọc đáp án từ bảng nếu có (bảo đảm 100% khớp bảng đáp án đính kèm)
        for tbl in self.doc.tables:
            for row in tbl.rows:
                for cell in row.cells:
                    ctext = cell.text.strip()
                    cell_m = re.search(r"(?:Câu\s+)?(\d+)[\.:\s]*([A-D])\b", ctext, re.IGNORECASE)
                    if cell_m:
                        c_num = int(cell_m.group(1))
                        c_key = cell_m.group(2).upper()
                        target_q = next((q for q in self.questions if q.original_index == c_num and q.part_type == PartType.PART_1_CHOICE), None)
                        if target_q and target_q.options:
                            for opt in target_q.options:
                                opt.is_correct = (opt.key.upper() == c_key)

        # Đọc đáp án Phần II Đúng/Sai từ bảng
        p2_questions = [q for q in self.questions if q.part_type == PartType.PART_2_TRUE_FALSE]
        p2_idx = 0
        for tbl in self.doc.tables:
            if len(tbl.rows) >= 2 and len(tbl.columns) == 4:
                headers = [c.text.strip().upper() for c in tbl.rows[0].cells]
                if headers == ['A', 'B', 'C', 'D']:
                    vals = [c.text.strip().upper() for c in tbl.rows[1].cells]
                    if all(v in ['Đ', 'S', 'D', 'S'] for v in vals) and p2_idx < len(p2_questions):
                        target_p2 = p2_questions[p2_idx]
                        for opt, val in zip(target_p2.options, vals):
                            opt.is_correct = (val in ['Đ', 'D'])
                        p2_idx += 1

        # Đọc đáp án Phần III từ bảng
        for tbl in self.doc.tables:
            if len(tbl.rows) >= 2:
                for c0, c1 in zip(tbl.rows[0].cells, tbl.rows[1].cells):
                    m = re.search(r"(?:Câu\s+)?(\d+)", c0.text.strip(), re.IGNORECASE)
                    if m:
                        q_num = int(m.group(1))
                        target_p3 = next((q for q in self.questions if q.original_index == q_num and q.part_type == PartType.PART_3_SHORT_ANSWER), None)
                        if target_p3:
                            ans_val = c1.text.strip()
                            if ans_val and not ans_val.lower().startswith("câu"):
                                target_p3.correct_short_answer = ans_val

        return self.questions

    def _process_question_block(self, block: List[Tuple[docx.text.paragraph.Paragraph, PartType, Optional[str]]], index: int) -> Optional[Question]:
        if not block:
            return None

        first_p, part_type, group_id = block[0]
        full_text = " ".join([p.text for p, _, _ in block])

        # Trích xuất Level [NB], [TH], [VD], [VDC]
        level = CognitiveLevel.UNKNOWN
        if re.search(r"\[NB\]", full_text, re.IGNORECASE):
            level = CognitiveLevel.NB
        elif re.search(r"\[TH\]", full_text, re.IGNORECASE):
            level = CognitiveLevel.TH
        elif re.search(r"\[VD\]", full_text, re.IGNORECASE):
            level = CognitiveLevel.VD
        elif re.search(r"\[VDC\]", full_text, re.IGNORECASE):
            level = CognitiveLevel.VDC

        # Trích xuất Chuyên đề [CĐ...]
        topic = ""
        topic_match = re.search(r"\[(CĐ\d+|Chủ đề[^\]]+)\]", full_text, re.IGNORECASE)
        if topic_match:
            topic = topic_match.group(1)

        # Khóa vị trí [KHÓA]
        is_locked_order = bool(re.search(r"\[KHÓA\]|\[KHOA\]", full_text, re.IGNORECASE))

        # Khóa đảo phương án [KĐ] hoặc '#'
        is_locked_options = bool(re.search(r"\[KĐ\]|\[KD\]|#", full_text, re.IGNORECASE))

        # Khởi tạo Question
        q = Question(
            id=f"Q_{index}",
            original_index=index,
            part_type=part_type,
            level=level,
            topic=topic,
            is_locked_order=is_locked_order,
            is_locked_options=is_locked_options,
            group_id=group_id,
        )

        # Tách Đề bài, Phương án, và Lời giải
        prompt_elements = []
        options = []
        solution_elements = []
        in_solution = False
        correct_short_answer = ""

        sol_start_pattern = re.compile(r"^\s*(Lời giải|Hướng dẫn giải)[\.:]?", re.IGNORECASE)
        short_ans_pattern = re.compile(r"^\s*Đáp án[\.:]?\s*([^\n\r]+)", re.IGNORECASE)

        # Xử lý các paragraph trong block
        for p, _, _ in block:
            p_text = p.text.strip()

            # Kiểm tra đáp án phần III (ưu tiên kiểm tra trước Lời giải)
            short_match = short_ans_pattern.match(p_text)
            if short_match:
                correct_short_answer = short_match.group(1).strip()
                continue

            if sol_start_pattern.match(p_text):
                in_solution = True
                solution_elements.append(copy.deepcopy(p._p))
                continue

            if in_solution:
                solution_elements.append(copy.deepcopy(p._p))
                continue

            # Phân tích phương án theo từng loại phần
            if part_type == PartType.PART_1_CHOICE:
                opts = self._extract_part1_options_from_paragraph(p)
                if opts:
                    options.extend(opts)
                else:
                    prompt_elements.append(copy.deepcopy(p._p))
            elif part_type == PartType.PART_2_TRUE_FALSE:
                opt = self._extract_part2_option_from_paragraph(p)
                if opt:
                    options.append(opt)
                else:
                    prompt_elements.append(copy.deepcopy(p._p))
            else:
                prompt_elements.append(copy.deepcopy(p._p))

        # Tự động phát hiện phương án logic phụ thuộc (Cả A và B, Tất cả các đáp án trên...)
        if part_type == PartType.PART_1_CHOICE and not q.is_locked_options:
            for opt in options:
                for regex in DEPENDENT_PHRASES:
                    if re.search(regex, opt.text_preview, re.IGNORECASE):
                        q.is_locked_options = True
                        self.warnings.append(
                            f"Câu {index}: Tự động BẬT KHÓA ĐẢO PHƯƠNG ÁN do phát hiện mệnh đề phụ thuộc '{opt.text_preview[:40]}...'"
                        )
                        break

        # Fix image-based options: Nếu Phần I không có text options nhưng có ảnh vẽ sẵn A, B, C, D
        if part_type == PartType.PART_1_CHOICE and not options:
            has_image = False
            for elem in prompt_elements:
                if any(c.tag.endswith("drawing") or c.tag.endswith("pict") for c in elem.iter()):
                    has_image = True
                    break
            if has_image:
                q.is_image_options = True
                q.is_locked_options = True
                options = [
                    Option(key="A", is_correct=False, text_preview="A"),
                    Option(key="B", is_correct=False, text_preview="B"),
                    Option(key="C", is_correct=False, text_preview="C"),
                    Option(key="D", is_correct=False, text_preview="D"),
                ] # Fixed image-based options

        q.prompt_xml_elements = prompt_elements
        q.options = options
        q.solution_xml_elements = solution_elements
        q.correct_short_answer = correct_short_answer
        q.raw_prompt_text = first_p.text[:120]
        return q

    def _extract_part1_options_from_paragraph(self, p: docx.text.paragraph.Paragraph) -> List[Option]:
        """Tách các phương án A, B, C, D từ paragraph và bảo tồn 100% các run MathType / OLE / OMML"""
        children = [c for c in p._p if not c.tag.endswith("}pPr")]
        child_texts = ["".join([t.text for t in c.iter() if t.tag.endswith("}t") and t.text]) for c in children]

        key_indices = []
        for i, txt in enumerate(child_texts):
            s = txt.strip()
            if re.match(r"^[A-D][\.:]", s):
                key = s[0]
                key_indices.append((i, key))
            elif s in ["A", "B", "C", "D"]:
                next_txt = "".join(child_texts[i+1:i+3]).strip()
                if next_txt.startswith(".") or next_txt.startswith(":"):
                    key_indices.append((i, s))

        if not key_indices:
            return []

        # Fix: Nếu câu dùng Word Numbered List (<w:numPr>), phương án A là text đầu trước B
        keys = [k for _, k in key_indices]
        if keys == ['B', 'C', 'D'] and key_indices[0][0] > 0:
            key_indices.insert(0, (0, 'A')) # Fixed numbered list missing A

        extracted = []
        for idx, (start, key) in enumerate(key_indices):
            end = key_indices[idx + 1][0] if idx + 1 < len(key_indices) else len(children)
            opt_runs = children[start:end]

            is_correct = False
            full_opt_text = "".join(child_texts[start:end]).strip()
            if f"*{key}" in full_opt_text:
                is_correct = True
            else:
                for r in opt_runs:
                    for u in r.iter():
                        if u.tag.endswith("}u") and u.attrib.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val") not in ["none", None]:
                            is_correct = True
                        if u.tag.endswith("}color") and u.attrib.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val") not in ["auto", "000000", None]:
                            is_correct = True

            text_preview = re.sub(r"^[A-D][\.:]\s*", "", full_opt_text).strip()

            opt = Option(
                key=key,
                is_correct=is_correct,
                text_preview=text_preview,
                xml_runs=[copy.deepcopy(r) for r in opt_runs]
            )
            extracted.append(opt)

        return extracted

    def _extract_part2_option_from_paragraph(self, p: docx.text.paragraph.Paragraph) -> Optional[Option]:
        """Tách các ý a), b), c), d) trong câu trắc nghiệm Đúng/Sai"""
        text = p.text.strip()
        m = re.match(r"^([a-d])\)\s*(.*)", text)
        if not m:
            return None

        key = m.group(1)
        content = m.group(2).strip()

        # Ý Đúng thường được gạch chân ở ký tự a), b)... hoặc tô đỏ (loại trừ màu đen 000000)
        is_correct = False
        for r in p.runs[:2]:
            has_underline = r.underline not in [False, None]
            has_color = False
            if r.font.color and r.font.color.rgb:
                hex_col = str(r.font.color.rgb).upper()
                if hex_col not in ["000000", "AUTO"]:
                    has_color = True
            if key in r.text and (has_underline or has_color):
                is_correct = True
                break

        return Option(
            key=key,
            is_correct=is_correct,
            text_preview=content,
            xml_runs=[copy.deepcopy(r._r) for r in p.runs]
        )
